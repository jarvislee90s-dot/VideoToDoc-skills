import tempfile
from pathlib import Path

import pytest

from videotodoc.document import ensure_mindmap_link


def test_ensure_mindmap_link_accepts_multiple_images():
    with tempfile.TemporaryDirectory() as tmp:
        md = Path(tmp) / "test.md"
        md.write_text("# Title\n\n## 图文讲义\n", encoding="utf-8")
        img1 = Path(tmp) / "mindmap_01.png"
        img2 = Path(tmp) / "mindmap_02.png"
        img1.write_text("", encoding="utf-8")
        img2.write_text("", encoding="utf-8")
        ensure_mindmap_link(md, [img1, img2])
        text = md.read_text(encoding="utf-8")
        assert "![思维导图](mindmap_01.png)" in text
        assert "![思维导图](mindmap_02.png)" in text


def test_ensure_mindmap_link_replaces_existing_single_image():
    with tempfile.TemporaryDirectory() as tmp:
        md = Path(tmp) / "test.md"
        md.write_text("# Title\n\n## 思维导图\n\n![思维导图](old.png)\n", encoding="utf-8")
        img = Path(tmp) / "new.png"
        img.write_text("", encoding="utf-8")
        ensure_mindmap_link(md, img)
        text = md.read_text(encoding="utf-8")
        assert "![思维导图](new.png)" in text
        assert "old.png" not in text


def test_ensure_mindmap_link_replaces_existing_multiple_images():
    with tempfile.TemporaryDirectory() as tmp:
        md = Path(tmp) / "test.md"
        md.write_text("# Title\n\n## 思维导图\n\n![思维导图](old_01.png)\n\n![思维导图](old_02.png)\n", encoding="utf-8")
        img1 = Path(tmp) / "new_01.png"
        img2 = Path(tmp) / "new_02.png"
        img1.write_text("", encoding="utf-8")
        img2.write_text("", encoding="utf-8")
        ensure_mindmap_link(md, [img1, img2])
        text = md.read_text(encoding="utf-8")
        assert "![思维导图](new_01.png)" in text
        assert "![思维导图](new_02.png)" in text
        assert "old_01.png" not in text
        assert "old_02.png" not in text


def test_generate_mindmap_adds_chapter_numbers(tmp_path: Path):
    from videotodoc.document import generate_mindmap
    from videotodoc.models import Section

    sections = [
        Section(slide_index=1, image_path="", start_ms=0, end_ms=1000, capture_ms=500, transcript="第一页内容", segment_indexes=[0]),
        Section(slide_index=2, image_path="", start_ms=1000, end_ms=2000, capture_ms=1500, transcript="第二页内容", segment_indexes=[1]),
    ]
    out = tmp_path / "mindmap.mmd"
    generate_mindmap("测试", sections, out)
    text = out.read_text(encoding="utf-8")
    assert "1. 第 1 页" in text
    assert "2. 第 2 页" in text


def test_generate_mindmap_keeps_root_syntax_valid(tmp_path: Path):
    from videotodoc.document import generate_mindmap
    from videotodoc.models import Section

    sections = [Section(slide_index=1, image_path="", start_ms=0, end_ms=1000, capture_ms=500, transcript="内容", segment_indexes=[0])]
    out = tmp_path / "mindmap.mmd"
    generate_mindmap("Git(Hub): 核心概念", sections, out)
    text = out.read_text(encoding="utf-8")
    assert "root((" in text
    assert "root（（" not in text


def test_process_result_mindmap_path_optional(tmp_path: Path):
    from videotodoc.models import ProcessResult
    result = ProcessResult(
        run_dir=tmp_path,
        transcript_path=tmp_path / "t.json",
        slides_path=tmp_path / "s.json",
        sections_path=tmp_path / "a.json",
        markdown_path=tmp_path / "m.md",
        mindmap_path=None,
    )
    assert result.mindmap_path is None


def test_markdown_to_docx_renders_markdown_table(tmp_path: Path):
    """Markdown 表格应渲染为真正的 Word 表格，而非残留 | 管道符原文。"""
    pytest.importorskip("docx", reason="python-docx 未安装，跳过 Word 渲染测试")
    from docx import Document
    from videotodoc.document import markdown_to_docx

    md = tmp_path / "t.md"
    md.write_text(
        "# 标题\n\n"
        "| 情景 | 定义 |\n"
        "|---|---|\n"
        "| **STEPS** | 已宣布政策情景 |\n"
        "| CPS | 当前政策情景 |\n",
        encoding="utf-8",
    )
    out = markdown_to_docx(md, tmp_path / "t.docx")
    doc = Document(str(out))

    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3
    assert table.rows[0].cells[0].text == "情景"
    assert table.rows[1].cells[0].text == "STEPS"
    # 表格内容不得以 | 原文形式残留在正文段落里
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "| 情景 |" not in body
    assert "---|" not in body


def test_markdown_to_docx_bold_runs_and_ordered_list(tmp_path: Path):
    """正文 **加粗** 应转为粗体 run（不留星号），有序列表应分行。"""
    pytest.importorskip("docx", reason="python-docx 未安装，跳过 Word 渲染测试")
    from docx import Document
    from videotodoc.document import markdown_to_docx

    md = tmp_path / "t.md"
    md.write_text(
        "# 标题\n\n"
        "增长 **20%** 至近 **1.2 万亿美元**。\n\n"
        "1. 第一类原型\n"
        "2. 第二类原型\n",
        encoding="utf-8",
    )
    out = markdown_to_docx(md, tmp_path / "t.docx")
    doc = Document(str(out))

    bold_texts = {
        run.text
        for para in doc.paragraphs
        for run in para.runs
        if run.bold
    }
    assert {"20%", "1.2 万亿美元"} <= bold_texts
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "**" not in all_text
    for item in ("第一类原型", "第二类原型"):
        assert any(item in p.text for p in doc.paragraphs)
